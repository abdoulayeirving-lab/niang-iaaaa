import csv
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from actions import data_analysis


class DataAnalysisTest(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.root = Path(self.temp_dir.name)

    def test_csv_statistics_handle_missing_and_categorical_values(self):
        path = self.root / "sales.csv"
        with path.open("w", encoding="utf-8", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(["Revenue", "Region"])
            writer.writerows([
                ["10", "North"],
                ["20", "South"],
                ["", "North"],
            ])

        rows, truncated, columns = data_analysis.analyze_dataset(path)

        self.assertEqual(rows, 3)
        self.assertFalse(truncated)
        self.assertEqual(columns[0]["kind"], "numeric")
        self.assertEqual(columns[0]["minimum"], "10")
        self.assertEqual(columns[0]["maximum"], "20")
        self.assertEqual(columns[0]["mean"], "15")
        self.assertEqual(columns[0]["missing"], 1)
        self.assertEqual(columns[1]["top_values"], "North (2), South (1)")

    def test_xlsx_is_analyzed_without_sending_dataset_to_ai(self):
        from openpyxl import Workbook

        path = self.root / "metrics.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Metric", "Value"])
        sheet.append(["CPU", 40])
        sheet.append(["RAM", 60])
        workbook.save(path)

        with (
            patch.object(data_analysis, "_resolve_source", return_value=str(path)),
            patch.object(data_analysis, "create_document", return_value="Document created: report.docx") as create_report,
        ):
            result = data_analysis.analyze_data({"source": str(path)})

        sections = create_report.call_args.kwargs["sections"]
        rows = sections[1]["table"]["rows"]
        self.assertEqual(rows[1][1], "numeric")
        self.assertIn("dataset contents were not uploaded", sections[0]["body"])
        self.assertIn("Document created: report.docx", result)

    def test_creates_a_real_word_report_from_csv(self):
        from docx import Document
        from actions import office_builder

        source = self.root / "sample.csv"
        source.write_text("Name,Score\nA,8\nB,10\n", encoding="utf-8")
        with patch.object(office_builder, "OUT_DIR", self.root / "reports"):
            result = data_analysis.analyze_data({"source": str(source)})

        self.assertTrue(result.startswith("Document created: "))
        report_path = Path(result.splitlines()[0].removeprefix("Document created: "))
        self.assertTrue(report_path.is_file())
        document = Document(report_path)
        self.assertEqual(len(document.tables), 1)
        self.assertEqual(document.tables[0].cell(1, 0).text, "Name")
        self.assertEqual(document.tables[0].cell(1, 4).text, "Top values: A (1), B (1)")

    def test_enforces_file_size_row_and_column_limits(self):
        path = self.root / "data.csv"
        path.write_text("a\n1\n2\n3\n", encoding="utf-8")
        with patch.object(data_analysis, "MAX_ROWS", 2):
            rows, truncated, _columns = data_analysis.analyze_dataset(path)
        self.assertEqual(rows, 2)
        self.assertTrue(truncated)

        too_many_columns = self.root / "wide.csv"
        too_many_columns.write_text(
            ",".join(f"c{i}" for i in range(data_analysis.MAX_COLUMNS + 1)),
            encoding="utf-8",
        )
        with self.assertRaisesRegex(ValueError, "columns"):
            data_analysis.analyze_dataset(too_many_columns)

        with (
            patch.object(data_analysis, "_resolve_source", return_value=str(path)),
            patch.object(data_analysis, "MAX_FILE_BYTES", 1),
        ):
            result = data_analysis.analyze_data({"source": str(path)})
        self.assertIn("too large", result)

    def test_rejects_non_tabular_and_remote_sources(self):
        with patch.object(data_analysis, "_resolve_source", return_value="https://example.com/data.csv"):
            result = data_analysis.analyze_data({"source": "https://example.com/data.csv"})
        self.assertIn("local CSV or Excel", result)

        path = self.root / "data.json"
        path.write_text("{}", encoding="utf-8")
        with patch.object(data_analysis, "_resolve_source", return_value=str(path)):
            result = data_analysis.analyze_data({"source": str(path)})
        self.assertIn("supports local .csv and .xlsx", result)


if __name__ == "__main__":
    unittest.main()
