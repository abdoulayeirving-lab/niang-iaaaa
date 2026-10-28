import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from actions import office_builder


class DocumentBuilderTest(unittest.TestCase):
    def test_creates_structured_docx(self):
        from docx import Document

        with TemporaryDirectory() as output:
            with patch.object(office_builder, "OUT_DIR", Path(output)):
                result = office_builder.create_document(
                    title="Project Report",
                    filename="../Unsafe:Report",
                    author="Test User",
                    sections=[
                        {
                            "title": "Summary",
                            "body": "First paragraph.\nSecond paragraph.",
                            "bullets": ["Finding one", "Finding two"],
                            "table": {
                                "headers": ["Metric", "Value"],
                                "rows": [["Users", "42"]],
                            },
                        }
                    ],
                )

            path = Path(result.removeprefix("Document created: "))
            self.assertTrue(path.is_file())
            self.assertEqual(path.parent, Path(output))
            self.assertEqual(path.name, "UnsafeReport.docx")
            document = Document(path)
            self.assertEqual(document.core_properties.title, "Project Report")
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("Summary", text)
            self.assertIn("First paragraph.", text)
            self.assertIn("Finding one", text)
            self.assertEqual(len(document.tables), 1)
            self.assertEqual(document.tables[0].cell(1, 1).text, "42")

    def test_rejects_missing_title_and_invalid_sections(self):
        self.assertIn("provide a document title", office_builder.create_document())
        self.assertIn(
            "must be an object",
            office_builder.create_document("Report", sections=["invalid"]),
        )

    def test_rejects_malformed_table(self):
        result = office_builder.create_document(
            "Report",
            sections=[{"title": "Data", "table": {"headers": [], "rows": []}}],
        )
        self.assertIn("requires at least one header", result)


if __name__ == "__main__":
    unittest.main()
