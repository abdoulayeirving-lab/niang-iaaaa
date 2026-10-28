import sys
from datetime import datetime
from pathlib import Path


def get_base_dir():
    if getattr(sys, "frozen", False):
        return Path(sys.executable).parent
    return Path(__file__).resolve().parent.parent


OUT_DIR = Path.home() / "Documents" / "Kaizumi"


def _ensure_out_dir() -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    return OUT_DIR


def _sanitize(name: str) -> str:
    name = "".join(c for c in str(name) if c not in '\\/:*?"<>|')
    name = name.strip().strip(". ")[:120].strip(". ")
    return name or "document"


def create_presentation(title: str = "", slides=None, filename: str = None) -> str:
    """Create a .pptx from a title and a list of slides.

    slides: list of dicts {"title": str, "bullets": [str]} or plain strings.
    """
    from pptx import Presentation

    if not title:
        return "Please provide a presentation title, sir."
    slides = slides or []

    try:
        prs = Presentation()
        title_layout = prs.slide_layouts[0]
        content_layout = prs.slide_layouts[1]

        s = prs.slides.add_slide(title_layout)
        s.shapes.title.text = title
        try:
            s.placeholders[1].text = datetime.now().strftime("%d.%m.%Y")
        except Exception:
            pass

        for idx, item in enumerate(slides, start=1):
            if isinstance(item, str):
                item = {"title": f"Slide {idx}", "bullets": [item]}
            slide_title = str(item.get("title") or f"Slide {idx}")
            bullets = item.get("bullets") or []
            if isinstance(bullets, str):
                bullets = [bullets]

            s = prs.slides.add_slide(content_layout)
            s.shapes.title.text = slide_title
            body = s.placeholders[1].text_frame
            first = True
            for b in bullets:
                p = body.paragraphs[0] if first else body.add_paragraph()
                first = False
                p.text = str(b)
                p.level = 0

        path = _ensure_out_dir() / f"{_sanitize(filename or title)}.pptx"
        prs.save(path)
        return f"Presentation created: {path}"
    except Exception as e:
        return f"Could not create presentation: {e}"


def create_document(
    title: str = "",
    sections=None,
    filename: str = None,
    author: str = "NIANG IA",
) -> str:
    """Create a structured Word document from sections, paragraphs, lists, and tables."""
    if not str(title).strip():
        return "Please provide a document title, sir."
    sections = sections or []
    if not isinstance(sections, list):
        return "Document sections must be a list."
    if len(sections) > 100:
        return "A document can contain at most 100 sections."
    for item in sections:
        if not isinstance(item, dict):
            return "Each document section must be an object."
        table = item.get("table")
        if table is not None:
            if not isinstance(table, dict):
                return "Section table must contain headers and rows."
            headers = table.get("headers") or []
            rows = table.get("rows") or []
            if not isinstance(headers, list) or not headers:
                return "A document table requires at least one header."
            if len(headers) > 50:
                return "A document table can contain at most 50 columns."
            if not isinstance(rows, list) or len(rows) > 5000:
                return "Document tables must have a list of at most 5,000 rows."
            if any(not isinstance(row, (list, tuple)) for row in rows):
                return "Each document table row must be a list."

    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt

        doc = Document()
        doc.core_properties.title = str(title).strip()
        doc.core_properties.author = str(author or "NIANG IA").strip()[:100]
        section = doc.sections[0]
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)

        heading = doc.add_heading(str(title).strip()[:250], level=0)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph(datetime.now().strftime("%d %B %Y")).alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        for item in sections:
            section_title = str(item.get("title") or "").strip()
            if section_title:
                try:
                    level = int(item.get("level", 1))
                except (TypeError, ValueError):
                    level = 1
                doc.add_heading(section_title[:250], level=max(1, min(level, 3)))

            body = item.get("body", "")
            if isinstance(body, str) and body.strip():
                for paragraph in body.split("\n"):
                    if paragraph.strip():
                        doc.add_paragraph(paragraph.strip())

            bullets = item.get("bullets") or []
            if isinstance(bullets, str):
                bullets = [bullets]
            if not isinstance(bullets, list):
                raise ValueError("Section bullets must be a list of strings.")
            for bullet in bullets:
                if str(bullet).strip():
                    doc.add_paragraph(str(bullet).strip(), style="List Bullet")

            table = item.get("table")
            if table is not None:
                headers = table.get("headers") or []
                rows = table.get("rows") or []
                word_table = doc.add_table(rows=1, cols=len(headers))
                word_table.style = "Light Shading Accent 1"
                for cell, value in zip(word_table.rows[0].cells, headers):
                    cell.text = str(value)
                for row in rows:
                    cells = word_table.add_row().cells
                    for index, cell in enumerate(cells):
                        cell.text = str(row[index]) if index < len(row) else ""

        for paragraph in doc.paragraphs:
            for run in paragraph.runs:
                if not run.font.name:
                    run.font.name = "Aptos"
                if run.font.size is None:
                    run.font.size = Pt(11)

        path = _ensure_out_dir() / f"{_sanitize(filename or title)}.docx"
        doc.save(path)
        return f"Document created: {path}"
    except Exception as e:
        return f"Could not create document: {e}"


def create_spreadsheet(
    filename: str = "spreadsheet",
    headers=None,
    rows=None,
    sheet_name: str = "Sheet1",
) -> str:
    """Create an .xlsx with optional headers and data rows."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    from openpyxl.utils import get_column_letter

    headers = headers or []
    rows = rows or []

    try:
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name

        if headers:
            for c, h in enumerate(headers, start=1):
                cell = ws.cell(row=1, column=c, value=str(h))
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center")

        start_row = 2 if headers else 1
        for r, row in enumerate(rows, start=start_row):
            if isinstance(row, str):
                row = [row]
            elif isinstance(row, (int, float)):
                row = [row]
            else:
                row = list(row)
            for c, v in enumerate(row, start=1):
                ws.cell(row=r, column=c, value=v)

        max_cols = max(len(headers), *(len(r) if not isinstance(r, (str, int, float)) else 1 for r in rows), 1)
        for c in range(1, max_cols + 1):
            letter = get_column_letter(c)
            ws.column_dimensions[letter].width = 14

        path = _ensure_out_dir() / f"{_sanitize(filename)}.xlsx"
        wb.save(path)
        return f"Spreadsheet created: {path}"
    except Exception as e:
        return f"Could not create spreadsheet: {e}"