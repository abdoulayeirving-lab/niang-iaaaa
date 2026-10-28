"""Bounded local analysis for CSV and Excel datasets."""

from __future__ import annotations

import csv
import math
import statistics
import zipfile
from collections import Counter
from pathlib import Path

from actions.document_qa import _resolve_source
from actions.office_builder import create_document
from openpyxl.utils.exceptions import InvalidFileException

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_ROWS = 10000
MAX_COLUMNS = 50
MAX_CATEGORIES = 5


def _read_csv(path: Path) -> tuple[list[str], list[list[object]], bool]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.reader(stream)
        headers = next(reader, None)
        if not headers:
            raise ValueError("The CSV file is empty.")
        if len(headers) > MAX_COLUMNS:
            raise ValueError(f"Dataset has more than {MAX_COLUMNS} columns.")
        rows = []
        truncated = False
        for row in reader:
            if len(rows) >= MAX_ROWS:
                truncated = True
                break
            rows.append(row[:MAX_COLUMNS])
        return [str(header)[:120] or f"Column {i + 1}" for i, header in enumerate(headers)], rows, truncated


def _read_xlsx(path: Path) -> tuple[list[str], list[list[object]], bool]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        iterator = worksheet.iter_rows(values_only=True)
        raw_headers = next(iterator, None)
        if not raw_headers:
            raise ValueError("The Excel sheet is empty.")
        raw_headers = list(raw_headers)
        if len(raw_headers) > MAX_COLUMNS:
            raise ValueError(f"Dataset has more than {MAX_COLUMNS} columns.")
        headers = [
            str(value)[:120] if value is not None and str(value).strip()
            else f"Column {index + 1}"
            for index, value in enumerate(raw_headers)
        ]
        rows = []
        truncated = False
        for row in iterator:
            if len(rows) >= MAX_ROWS:
                truncated = True
                break
            rows.append(list(row[:MAX_COLUMNS]))
        return headers, rows, truncated
    finally:
        workbook.close()


def _number(value: object) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        number = float(value)
        return number if math.isfinite(number) else None
    text = str(value).strip().replace(",", "")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _format_number(value: float) -> str:
    return f"{value:,.2f}".rstrip("0").rstrip(".")


def analyze_dataset(path: Path) -> tuple[int, bool, list[dict]]:
    if path.suffix.lower() == ".csv":
        headers, rows, truncated = _read_csv(path)
    elif path.suffix.lower() == ".xlsx":
        headers, rows, truncated = _read_xlsx(path)
    else:
        raise ValueError("Only .csv and .xlsx datasets are supported.")
    if not headers:
        raise ValueError("The dataset has no columns.")

    analyses = []
    for column_index, header in enumerate(headers):
        values = [
            row[column_index] if column_index < len(row) else None
            for row in rows
        ]
        present = [
            value for value in values
            if value is not None and str(value).strip()
        ]
        numeric = [
            number for value in present
            if (number := _number(value)) is not None
        ]
        summary = {
            "column": header,
            "non_empty": len(present),
            "missing": len(values) - len(present),
        }
        if numeric and len(numeric) == len(present):
            summary["kind"] = "numeric"
            summary["minimum"] = _format_number(min(numeric))
            summary["maximum"] = _format_number(max(numeric))
            summary["mean"] = _format_number(statistics.fmean(numeric))
            summary["median"] = _format_number(statistics.median(numeric))
        else:
            summary["kind"] = "categorical"
            counter = Counter(str(value).strip()[:100] for value in present)
            summary["top_values"] = ", ".join(
                f"{label} ({count})"
                for label, count in counter.most_common(MAX_CATEGORIES)
            ) or "No values"
        analyses.append(summary)
    return len(rows), truncated, analyses


def analyze_data(
    parameters: dict,
    response=None,
    player=None,
    session_memory=None,
) -> str:
    """Analyze a local CSV/XLSX and create a descriptive Word report."""
    params = parameters or {}
    source = _resolve_source(str(params.get("source", "")).strip())
    if not source or source.lower().startswith(("http://", "https://")):
        return "Give me a local CSV or Excel file path to analyze, sir."

    path = Path(source)
    if path.suffix.lower() not in {".csv", ".xlsx"}:
        return "Data analysis currently supports local .csv and .xlsx files."
    try:
        if path.stat().st_size > MAX_FILE_BYTES:
            return f"Dataset is too large; the limit is {MAX_FILE_BYTES // (1024 * 1024)} MB."
        row_count, truncated, analyses = analyze_dataset(path)
    except (OSError, ValueError, UnicodeError, csv.Error, zipfile.BadZipFile, InvalidFileException) as exc:
        return f"Could not analyze dataset: {exc}"

    summary_rows = []
    for item in analyses:
        if item["kind"] == "numeric":
            details = (
                f"min={item['minimum']}; max={item['maximum']}; "
                f"mean={item['mean']}; median={item['median']}"
            )
        else:
            details = f"Top values: {item['top_values']}"
        summary_rows.append([
            item["column"],
            item["kind"],
            str(item["non_empty"]),
            str(item["missing"]),
            details,
        ])

    title = f"Data analysis - {path.stem}"
    limitations = (
        f"Analyzed {row_count} data row(s) across {len(analyses)} column(s). "
        "Numeric statistics and category counts were computed locally; dataset "
        "contents were not uploaded to an AI service."
    )
    if truncated:
        limitations += (
            f" The report is based on the first {MAX_ROWS} rows; the dataset "
            "contains additional rows that were not analyzed."
        )

    report_result = create_document(
        title=title,
        filename=str(params.get("report_name") or title),
        sections=[
            {"title": "Overview", "body": limitations},
            {
                "title": "Column analysis",
                "body": (
                    "Missing counts are based on the analyzed rows. "
                    "Numeric columns include min, max, mean, and median. "
                    "Categorical columns list the most frequent values."
                ),
                "table": {
                    "headers": ["Column", "Type", "Non-empty", "Missing", "Summary"],
                    "rows": summary_rows,
                },
            },
        ],
    )
    if not report_result.startswith("Document created: "):
        return report_result
    return f"{report_result}\n{limitations}"
